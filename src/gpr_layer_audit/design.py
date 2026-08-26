from __future__ import annotations

import csv
from dataclasses import replace
from pathlib import Path

from openpyxl import load_workbook

from gpr_layer_audit.models import DesignSegment, ThicknessResult

DESIGN_COLUMNS = {
    "road_id",
    "start_chainage_m",
    "end_chainage_m",
    "layer_name",
    "design_thickness_mm",
}


def read_design_schedule(path: str | Path) -> list[DesignSegment]:
    file_path = Path(path)
    if file_path.suffix.lower() == ".csv":
        with file_path.open("r", encoding="utf-8-sig", newline="") as stream:
            rows = list(csv.DictReader(stream))
    elif file_path.suffix.lower() in {".xlsx", ".xlsm"}:
        workbook = load_workbook(file_path, read_only=True, data_only=True)
        sheet = workbook.active
        headers = [str(cell.value or "").strip() for cell in next(sheet.iter_rows())]
        rows = [
            dict(zip(headers, (cell.value for cell in row), strict=False))
            for row in sheet.iter_rows()
        ]
        rows = rows[1:]
    else:
        raise ValueError("Design schedule must be CSV or XLSX")
    if rows and not DESIGN_COLUMNS.issubset(rows[0]):
        missing = ", ".join(sorted(DESIGN_COLUMNS - set(rows[0])))
        raise ValueError(f"Design schedule is missing required columns: {missing}")
    output: list[DesignSegment] = []
    for row in rows:
        if not row or row.get("start_chainage_m") in {None, ""}:
            continue
        output.append(
            DesignSegment(
                road_id=str(row["road_id"]),
                start_chainage_m=float(row["start_chainage_m"]),
                end_chainage_m=float(row["end_chainage_m"]),
                layer_name=str(row["layer_name"]),
                design_thickness_mm=float(row["design_thickness_mm"]),
                tolerance_low_mm=(
                    float(row["tolerance_low_mm"])
                    if row.get("tolerance_low_mm") not in {None, ""}
                    else None
                ),
                tolerance_high_mm=(
                    float(row["tolerance_high_mm"])
                    if row.get("tolerance_high_mm") not in {None, ""}
                    else None
                ),
            )
        )
    return output


def compare_with_design(
    results: list[ThicknessResult], segments: list[DesignSegment]
) -> list[ThicknessResult]:
    compared: list[ThicknessResult] = []
    for item in results:
        match = next(
            (
                segment
                for segment in segments
                if segment.layer_name.casefold() == item.layer_name.casefold()
                and segment.start_chainage_m <= item.chainage_m <= segment.end_chainage_m
            ),
            None,
        )
        if match is None or item.thickness_mm is None:
            compared.append(item)
            continue
        deviation = item.thickness_mm - match.design_thickness_mm
        percent = deviation / match.design_thickness_mm * 100 if match.design_thickness_mm else None
        compliance = "not_evaluated"
        if match.tolerance_low_mm is not None or match.tolerance_high_mm is not None:
            lower = match.tolerance_low_mm if match.tolerance_low_mm is not None else float("-inf")
            upper = match.tolerance_high_mm if match.tolerance_high_mm is not None else float("inf")
            compliance = "compliant" if lower <= deviation <= upper else "out_of_tolerance"
        compared.append(
            replace(
                item,
                design_thickness_mm=match.design_thickness_mm,
                deviation_mm=deviation,
                deviation_percent=percent,
                compliance=compliance,
            )
        )
    return compared
