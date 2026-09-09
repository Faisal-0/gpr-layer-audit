from __future__ import annotations

import csv
import json
import shutil
from dataclasses import asdict
from datetime import UTC, datetime
from pathlib import Path

import matplotlib
import numpy as np
from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from gpr_layer_audit import __version__
from gpr_layer_audit.models import AnalysisResult, InterfacePick, PickStatus, VisibilityState
from gpr_layer_audit.seeds import seed_document
from gpr_layer_audit.time_coordinates import image_time_bounds_ns, sample_time_ns

matplotlib.use("Agg")
from matplotlib import pyplot as plt  # noqa: E402


def _serialise(value):
    if hasattr(value, "value"):
        return value.value
    if isinstance(value, Path):
        return str(value)
    if isinstance(value, np.generic):
        return value.item()
    raise TypeError(type(value).__name__)


def _rows(items: list) -> tuple[list[str], list[list]]:
    if not items:
        return [], []
    dictionaries = [asdict(item) for item in items]
    headers = list(dictionaries[0])

    def cell(value):
        if hasattr(value, "value"):
            return value.value
        if isinstance(value, (list, dict, tuple)):
            return json.dumps(value, default=_serialise)
        return value

    return headers, [[cell(value) for value in row.values()] for row in dictionaries]


def _sheet(workbook: Workbook, title: str, headers: list[str], rows: list[list]) -> None:
    sheet = workbook.create_sheet(title)
    if headers:
        sheet.append(headers)
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="203344")
            cell.alignment = Alignment(vertical="center")
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(1, len(rows) + 1)}"
    for row in rows:
        sheet.append(row)
    for column, header in enumerate(headers, 1):
        values = [str(header), *(str(row[column - 1] or "") for row in rows[:300])]
        sheet.column_dimensions[get_column_letter(column)].width = min(
            34, max(10, max(map(len, values)) + 2)
        )


def _workbook(result: AnalysisResult, path: Path) -> None:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    summary.append(["GPR Layer Audit", "Physics-first pavement thickness analysis"])
    summary.append(["Source", str(result.source.dzt_path)])
    summary.append(["Source SHA-256", result.source.fingerprint])
    summary.append(["Antenna", result.header.antenna])
    summary.append(["Traces", result.header.trace_count])
    summary.append(["Samples per trace", result.header.samples_per_trace])
    summary.append(["Time range (ns)", result.header.range_ns])
    summary.append(["Stack size", result.stack_size])
    summary.append(["Plate valid for dielectric", result.diagnostics.valid_for_dielectric])
    summary.append(["Calibration notes", " | ".join(result.diagnostics.messages)])
    summary.column_dimensions["A"].width = 30
    summary.column_dimensions["B"].width = 90
    summary["A1"].font = Font(bold=True, size=16, color="FFFFFF")
    summary["B1"].font = Font(bold=True, color="FFFFFF")
    for cell in summary[1]:
        cell.fill = PatternFill("solid", fgColor="203344")

    thickness_headers, thickness_rows = _rows(result.thickness)
    _sheet(workbook, "Thickness Results", thickness_headers, thickness_rows)
    design_rows = (
        [
            row
            for row in thickness_rows
            if row[thickness_headers.index("design_thickness_mm")] is not None
        ]
        if thickness_headers
        else []
    )
    _sheet(workbook, "Design Comparison", thickness_headers, design_rows)
    issue_headers, issue_rows = _rows(result.review_issues)
    _sheet(workbook, "Exceptions", issue_headers, issue_rows)
    pick_headers, pick_rows = _rows(result.picks)
    _sheet(workbook, "Interface Picks", pick_headers, pick_rows)
    reference_headers, reference_rows = _rows(result.reference_diagnostics)
    _sheet(workbook, "Manual Reference Diagnostic", reference_headers, reference_rows)
    seed_rows = [
        [
            item.station_id,
            item.chainage_m,
            item.role,
            json.dumps(item.samples, default=_serialise),
            json.dumps(item.visibility, default=_serialise),
        ]
        for item in result.seed_stations
    ]
    _sheet(
        workbook,
        "Seed Stations",
        ["station_id", "chainage_m", "role", "samples", "visibility"],
        seed_rows,
    )
    profile_headers, profile_rows = _rows(result.profile)
    _sheet(workbook, "Layer Profiles", profile_headers, profile_rows)
    anomaly_headers, anomaly_rows = _rows(result.anomaly_regions)
    _sheet(workbook, "Structural Anomalies", anomaly_headers, anomaly_rows)
    candidate_headers, candidate_rows = _rows(result.candidate_events)
    _sheet(workbook, "Candidate Events", candidate_headers, candidate_rows)
    retention_headers, retention_rows = _rows(result.retention_audit)
    _sheet(workbook, "Retention Audit", retention_headers, retention_rows)
    method = workbook.create_sheet("Method & Provenance")
    method.append(["Software version", __version__])
    method.append(["Generated UTC", datetime.now(UTC).isoformat()])
    method.append(["Parameters", json.dumps(result.parameters, default=_serialise, indent=2)])
    method.append(
        [
            "Interpretation notice",
            "Thickness is estimated from TWTT and the recorded dielectric source; "
            "it is not independently verified core truth.",
        ]
    )
    method.column_dimensions["A"].width = 28
    method.column_dimensions["B"].width = 110
    if "Thickness Results" in workbook.sheetnames and thickness_headers:
        sheet = workbook["Thickness Results"]
        status_col = thickness_headers.index("status") + 1
        col = get_column_letter(status_col)
        sheet.conditional_formatting.add(
            f"{col}2:{col}{sheet.max_row}",
            CellIsRule(
                operator="equal", formula=['"review"'], fill=PatternFill("solid", fgColor="FFF2CC")
            ),
        )
    workbook.save(path)


def _csv(result: AnalysisResult, path: Path) -> None:
    headers, rows = _rows(result.thickness)
    with path.open("w", newline="", encoding="utf-8-sig") as stream:
        writer = csv.writer(stream)
        writer.writerow(headers)
        writer.writerows(rows)


def _write_rows_csv(items: list, path: Path) -> None:
    headers, rows = _rows(items)
    with path.open("w", newline="", encoding="utf-8-sig") as stream:
        writer = csv.writer(stream)
        writer.writerow(headers)
        writer.writerows(rows)


def _geojson(result: AnalysisResult, path: Path) -> None:
    features = []
    for item in result.thickness:
        if item.latitude is None or item.longitude is None:
            continue
        properties = asdict(item)
        properties.pop("latitude", None)
        properties.pop("longitude", None)
        features.append(
            {
                "type": "Feature",
                "geometry": {"type": "Point", "coordinates": [item.longitude, item.latitude]},
                "properties": properties,
            }
        )
    path.write_text(
        json.dumps(
            {"type": "FeatureCollection", "features": features}, default=_serialise, indent=2
        ),
        encoding="utf-8",
    )


def _overlay_series(items: list[InterfacePick], statuses: set[PickStatus]) -> np.ndarray:
    """Retain missing rows as NaN: filtering them would draw through evidence gaps."""
    return np.asarray(
        [
            (
                item.selected_lobe_sample
                if item.selected_lobe_sample is not None
                else item.sample_index
            )
            if item.sample_index >= 0
            and item.status in statuses
            and not item.anomaly
            and item.visibility not in {VisibilityState.ABSENT, VisibilityState.NOT_VISIBLE}
            else np.nan
            for item in items
        ]
    )


def _radargram(
    result: AnalysisResult,
    path: Path,
    start_chainage_m: float | None = None,
    end_chainage_m: float | None = None,
    title: str = "Calibrated radargram and interpreted interfaces",
    view_name: str = "Clean",
) -> None:
    figure, axis = plt.subplots(figsize=(16, 7), constrained_layout=True)
    start = float(result.chainage_m[0]) if start_chainage_m is None else start_chainage_m
    end = float(result.chainage_m[-1]) if end_chainage_m is None else end_chainage_m
    selected = (result.chainage_m >= start) & (result.chainage_m <= end)
    indices = np.flatnonzero(selected)
    if not len(indices):
        plt.close(figure)
        return
    source = result.display_radargrams.get(view_name, result.calibrated_radargram)
    data = source[indices].T
    limit = float(np.percentile(np.abs(data), 98.5)) or 1.0
    top_time, bottom_time = image_time_bounds_ns(result)
    extent = [
        float(result.chainage_m[indices[0]]),
        float(result.chainage_m[indices[-1]]),
        bottom_time,
        top_time,
    ]
    if result.parameters.get("input_mode") == "processed":
        step = result.header.distance_per_trace_m * result.parameters["processed_stride"]
        extent[0] -= step / 2
        extent[1] += step / 2
    axis.imshow(data, cmap="gray", aspect="auto", vmin=-limit, vmax=limit, extent=extent)
    colours = {1: "#28d7e5", 2: "#ffc857", 3: "#ff6b6b"}
    for order in sorted({item.layer_order for item in result.picks}):
        items = [
            item
            for item in result.picks
            if item.layer_order == order and start <= item.chainage_m <= end
        ]
        items.sort(key=lambda item: item.chainage_m)
        if not items:
            continue
        for statuses, style, label in (
            ({PickStatus.HIGH_CONFIDENCE, PickStatus.ACCEPTED}, "-", "accepted"),
            ({PickStatus.REVIEW}, ":", "review candidate"),
        ):
            values = sample_time_ns(result, _overlay_series(items, statuses))
            if np.any(np.isfinite(values)):
                axis.plot(
                    [item.chainage_m for item in items],
                    values,
                    color=colours.get(order, "#ffffff"),
                    linewidth=1.1,
                    linestyle=style,
                    label=f"{items[0].layer_name} ({label})",
                )
    axis.set_xlabel("Chainage (m)")
    axis.set_ylabel("Time (ns)")
    axis.set_title(f"{title} · {view_name} view")
    axis.legend(loc="upper right")
    figure.savefig(path, dpi=180)
    plt.close(figure)


def _profile_png(result: AnalysisResult, path: Path) -> None:
    figure, axes = plt.subplots(2, 1, figsize=(16, 8), sharex=True, constrained_layout=True)
    colours = {1: "#28d7e5", 2: "#ffc857", 3: "#ff6b6b"}
    for order in sorted({item.layer_order for item in result.profile}):
        points = sorted(
            (item for item in result.profile if item.layer_order == order),
            key=lambda item: item.chainage_m,
        )
        x = np.asarray([item.chainage_m for item in points])
        cumulative = np.asarray(
            [
                item.cumulative_depth_mm if item.cumulative_depth_mm is not None else np.nan
                for item in points
            ]
        )
        individual = np.asarray(
            [
                item.individual_thickness_mm if item.individual_thickness_mm is not None else np.nan
                for item in points
            ]
        )
        axes[0].plot(x, cumulative, color=colours.get(order), label=points[0].layer_name)
        axes[1].plot(x, individual, color=colours.get(order), label=points[0].layer_name)
        design = np.asarray(
            [
                item.design_thickness_mm if item.design_thickness_mm is not None else np.nan
                for item in points
            ]
        )
        if np.any(np.isfinite(design)):
            axes[1].plot(x, design, color=colours.get(order), linestyle="--", alpha=0.65)
    for region in result.anomaly_regions:
        for axis in axes:
            axis.axvspan(
                region.start_chainage_m, region.end_chainage_m, color="#ff6b6b", alpha=0.16
            )
    axes[0].set_ylabel("Cumulative interface depth (mm)")
    axes[1].set_ylabel("Individual thickness (mm)")
    axes[1].set_xlabel("Chainage (m)")
    for axis in axes:
        axis.grid(alpha=0.2)
        axis.legend(loc="upper right")
    figure.savefig(path, dpi=180)
    plt.close(figure)


def export_audit_package(result: AnalysisResult, output_directory: str | Path) -> Path:
    root = Path(output_directory)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    package = root / f"gpr_audit_{timestamp}"
    package.mkdir(parents=True, exist_ok=False)
    _workbook(result, package / "gpr_layer_audit.xlsx")
    _csv(result, package / "thickness_results.csv")
    _write_rows_csv(result.picks, package / "interface_observations.csv")
    _write_rows_csv(result.profile, package / "layer_profiles.csv")
    _write_rows_csv(result.retention_audit, package / "retention_audit.csv")
    _geojson(result, package / "thickness_results.geojson")
    _radargram(result, package / "annotated_radargram.png")
    _profile_png(result, package / "layer_profiles.png")
    views = package / "processing_views"
    views.mkdir()
    for view_name in ("Raw", "Clean", "Phase", "Gradient", "Candidates"):
        if view_name in result.display_radargrams:
            _radargram(
                result,
                views / f"{view_name.casefold()}_overview.png",
                title="Interpretation processing comparison",
                view_name=view_name,
            )
    radargrams = package / "radargrams"
    radargrams.mkdir()
    _radargram(
        result,
        radargrams / "accepted_overview.png",
        title="Accepted measurements and review overlays",
    )
    exception_directory = radargrams / "exceptions"
    exception_directory.mkdir()
    for index, issue in enumerate(result.review_issues, 1):
        margin = max(5.0, (issue.end_chainage_m - issue.start_chainage_m) * 0.1)
        _radargram(
            result,
            exception_directory / f"exception_{index:03d}_layer_{issue.layer_order}.png",
            max(0.0, issue.start_chainage_m - margin),
            issue.end_chainage_m + margin,
            f"Exception {index}: {issue.layer_name} "
            f"{issue.start_chainage_m:.1f}–{issue.end_chainage_m:.1f} m",
        )
    manifest = result.manifest()
    if result.sample_validity is not None:
        np.savez_compressed(
            package / "signal_evidence.npz",
            sample_validity=result.sample_validity,
            trace_chainage_m=result.chainage_m,
            surface_samples_raw=result.surface_samples_raw,
        )
        manifest["signal_evidence"] = {
            "file": "signal_evidence.npz",
            "axes": ["trace", "sample"],
            "invalid_samples": int(np.count_nonzero(~result.sample_validity)),
        }
    layer_orders = sorted({item.layer_order for item in result.picks})
    manifest["review_coverage"] = {
        str(order): (
            sum(
                not item.is_accepted_measurement
                for item in result.picks
                if item.layer_order == order
            )
            / max(1, sum(item.layer_order == order for item in result.picks))
        )
        for order in layer_orders
    }
    manifest.update(
        {"software_version": __version__, "generated_utc": datetime.now(UTC).isoformat()}
    )
    (package / "manifest.json").write_text(
        json.dumps(manifest, default=_serialise, indent=2), encoding="utf-8"
    )
    (package / "seeds.json").write_text(
        json.dumps(
            seed_document(
                str(result.parameters.get("survey_id") or result.source.dzt_path.stem),
                result.seed_stations,
                layer_names={item.layer_order: item.layer_name for item in result.picks},
            ),
            indent=2,
        ),
        encoding="utf-8",
    )
    shutil.make_archive(str(package), "zip", root_dir=package)
    return package
