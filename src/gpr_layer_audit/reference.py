from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook

from gpr_layer_audit.models import (
    AnalysisResult,
    DielectricSource,
    PickStatus,
    ReferenceDiagnostic,
)
from gpr_layer_audit.processing.dielectric import thickness_from_twtt_mm


@dataclass(slots=True)
class ManualReferencePoint:
    layer_order: int
    chainage_m: float
    interface_depth_mm: float
    interpolated: bool = False


def _is_interpolated(cell) -> bool:
    colour = cell.fill.fgColor
    return cell.fill.fill_type == "solid" and str(colour.rgb).upper().endswith("FFF2CC")


def read_manual_reference(path: str | Path) -> list[ManualReferencePoint]:
    workbook = load_workbook(path, data_only=True, read_only=False)
    sheet = (
        workbook["Interpolated Data"]
        if "Interpolated Data" in workbook.sheetnames
        else workbook.active
    )
    headers = {str(cell.value or "").strip().lower(): cell.column for cell in sheet[1]}
    distance_column = next(
        (column for name, column in headers.items() if name.startswith("dist")), None
    )
    if distance_column is None:
        raise ValueError("Manual reference workbook needs a distance/chainage column.")
    layer_columns = sorted(
        (
            int(name.split("layer", 1)[1].split()[0]),
            column,
            "(in)" in name or "inch" in name,
        )
        for name, column in headers.items()
        if name.startswith("layer") and "depth" in name
    )
    if not layer_columns:
        raise ValueError("Manual reference workbook has no 'Layer N Depth' columns.")
    output: list[ManualReferencePoint] = []
    for row in range(2, sheet.max_row + 1):
        distance = sheet.cell(row, distance_column).value
        if distance is None:
            continue
        for order, column, inches in layer_columns:
            cell = sheet.cell(row, column)
            if cell.value is None:
                continue
            depth_mm = float(cell.value) * 25.4 if inches else float(cell.value)
            output.append(
                ManualReferencePoint(
                    layer_order=order,
                    chainage_m=float(distance),
                    interface_depth_mm=depth_mm,
                    interpolated=_is_interpolated(cell),
                )
            )
    return output


def evaluate_manual_reference(
    result: AnalysisResult,
    points: list[ManualReferencePoint],
) -> list[ReferenceDiagnostic]:
    picks_by_layer = {
        order: [item for item in result.picks if item.layer_order == order]
        for order in {item.layer_order for item in result.picks}
    }
    dielectric_parameters = result.parameters.get("dielectric_by_layer", {})
    output: list[ReferenceDiagnostic] = []
    for point in points:
        cumulative_depth = 0.0
        previous_sample = float(result.reference_surface_sample)
        selected_status = PickStatus.UNRESOLVED
        measurable = True
        for order in range(1, point.layer_order + 1):
            candidates = picks_by_layer.get(order, [])
            if not candidates:
                measurable = False
                break
            selected = min(candidates, key=lambda item: abs(item.chainage_m - point.chainage_m))
            selected_status = selected.status
            dielectric_item = dielectric_parameters.get(order) or dielectric_parameters.get(
                str(order)
            )
            dielectric = dielectric_item.get("value") if dielectric_item else None
            source = dielectric_item.get("source") if dielectric_item else None
            if dielectric is None or source == DielectricSource.UNRESOLVED:
                measurable = False
                break
            twtt = (
                max(0.0, selected.sample_index - previous_sample) * result.header.sample_interval_ns
            )
            cumulative_depth += thickness_from_twtt_mm(twtt, float(dielectric))
            previous_sample = selected.sample_index
        measured = cumulative_depth if measurable else None
        error = abs(measured - point.interface_depth_mm) if measured is not None else None
        target = 12.7 if point.layer_order == 1 else 25.4
        output.append(
            ReferenceDiagnostic(
                layer_order=point.layer_order,
                chainage_m=point.chainage_m,
                reference_interface_depth_mm=point.interface_depth_mm,
                measured_interface_depth_mm=measured,
                absolute_error_mm=error,
                interpolated_reference=point.interpolated,
                pick_status=selected_status,
                within_release_target=error <= target if error is not None else None,
            )
        )
    result.reference_diagnostics = output
    return output
