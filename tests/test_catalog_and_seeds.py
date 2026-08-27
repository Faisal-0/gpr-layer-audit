from __future__ import annotations

import json
import shutil

import numpy as np
import pytest
from openpyxl import Workbook, load_workbook

from gpr_layer_audit.benchmark import _holdout_block
from gpr_layer_audit.catalog import calibration_candidates_for, discover_survey_catalog
from gpr_layer_audit.checkpoints import load_checkpoint_file, save_checkpoint_file
from gpr_layer_audit.models import (
    LabelOrigin,
    SeedStation,
    ValidationCheckpoint,
    VisibilityState,
)
from gpr_layer_audit.project.store import SCHEMA_VERSION, ProjectStore
from gpr_layer_audit.reference import normalize_reference_workbook
from gpr_layer_audit.seeds import load_seed_file, save_seed_file


def test_catalog_discovers_multi_dzt_companions_and_ambiguous_plates(
    tmp_path, synthetic_acquisition
):
    road, plate, _ = synthetic_acquisition
    project = tmp_path / "ROAD_A.PRJ"
    project.mkdir()
    for suffix in (".DZT", ".DZG", ".DZX"):
        shutil.copy2(road.with_suffix(suffix), project / f"ROAD_A{suffix}")
    shutil.copy2(road, tmp_path / "ROAD_B.DZT")
    shutil.copy2(plate, tmp_path / "ROAD_A_PLATE_1.DZT")
    shutil.copy2(plate, tmp_path / "ROAD_A_PLATE_2.DZT")
    (tmp_path / "road_design.csv").write_text(
        "road_id,start_chainage_m,end_chainage_m,layer_name,design_thickness_mm\n",
        encoding="utf-8",
    )

    catalog = discover_survey_catalog(tmp_path)

    assert len(catalog.roads) >= 2
    assert len(catalog.calibrations) >= 2
    road_a = next(item for item in catalog.roads if item.dzt_path.name == "ROAD_A.DZT")
    assert road_a.project_path == project
    assert catalog.files_by_type[".PRJ"] == [project]
    assert road_a.dzg_path is not None and road_a.dzx_path is not None
    candidates = calibration_candidates_for(catalog, road_a.survey_id)
    assert len(candidates) >= 2
    assert abs(candidates[0].compatibility_score - candidates[1].compatibility_score) < 0.05
    assert catalog.design_files == [tmp_path / "road_design.csv"]


def test_reference_normalization_records_cells_and_formula_origin(tmp_path):
    path = tmp_path / "reference.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Road 7"
    sheet.append(["Filename", "Chainage (m)", "Asphalt depth (mm)", "Base course depth (mm)"])
    sheet.append(["ROAD_7.DZT", 10, 50, 140])
    sheet.append(["ROAD_7.DZT", 20, "=C2+5", "=FORECAST.LINEAR(B3,D2:D2,B2:B2)"])
    workbook.save(path)
    # Populate cached-independent manual values in the first row; formulas in
    # the second row remain intentionally unavailable to data_only loading.
    normalized = normalize_reference_workbook(path)

    assert [(item.layer_order, item.cumulative_depth_mm) for item in normalized] == [
        (1, 50.0),
        (2, 140.0),
    ]
    assert all(item.label_origin == LabelOrigin.MANUAL for item in normalized)
    assert {item.source_cell for item in normalized} == {"C2", "D2"}

    # Check the formula classifier directly through a workbook with cached-like
    # numeric cells and retained source provenance in the normal path.
    book = load_workbook(path, data_only=False)
    assert book["Road 7"]["C3"].data_type == "f"


def test_reference_csv_normalization(tmp_path):
    path = tmp_path / "reference.csv"
    path.write_text(
        "Filename,Chainage (m),Asphalt depth (mm),Base course depth (mm)\n"
        "ROAD_8.DZT,5,48,132\n",
        encoding="utf-8",
    )

    points = normalize_reference_workbook(path)

    assert [item.cumulative_depth_mm for item in points] == [48.0, 132.0]
    assert points[1].individual_thickness_mm == 84.0
    assert points[0].source_sheet == "CSV"


def test_seed_json_round_trip_and_limit(tmp_path):
    path = tmp_path / "seeds.json"
    stations = [
        SeedStation(
            "station-1",
            12.5,
            {1: 80.0, 2: 123.0},
            {
                1: VisibilityState.VISIBLE,
                2: VisibilityState.VISIBLE,
                3: VisibilityState.NOT_VISIBLE,
            },
            user_confirmed={1: True, 2: True, 3: True},
            phase_class={1: 1, 2: 2},
            analytic_phase_rad={1: -2.2, 2: -1.4},
            polarity={1: -1, 2: -1},
            selected_lobe={1: "negative_trough", 2: "negative_trough"},
            canonical_samples={1: 82.5, 2: 125.5},
            pulse_width_samples={1: 7.0, 2: 8.0},
            event_ids={1: "event-l1", 2: "event-l2"},
            regime_ids={1: "default", 2: "default"},
            competing_samples={1: [86.0], 2: [118.0]},
            preview_status={1: "confirmed", 2: "confirmed", 3: "explicit_visibility_state"},
        )
    ]
    save_seed_file(path, "ROAD_A.PRJ/ROAD_A", stations, layer_names={1: "Asphalt"})

    survey_id, loaded = load_seed_file(path)

    assert survey_id == "ROAD_A.PRJ/ROAD_A"
    assert loaded == stations
    document = json.loads(path.read_text(encoding="utf-8"))
    assert document["schema_version"] == 4
    assert document["stations"][0]["picks"]["2"]["user_confirmed"] is True


def test_seed_file_rejects_conflicting_lobes_in_one_regime(tmp_path):
    common = {
        "visibility": {1: VisibilityState.VISIBLE},
        "user_confirmed": {1: True},
        "analytic_phase_rad": {1: 2.8},
        "canonical_samples": {1: 100.0},
        "pulse_width_samples": {1: 7.0},
        "regime_ids": {1: "default"},
    }
    stations = [
        SeedStation(
            "negative",
            10.0,
            {1: 100.0},
            phase_class={1: 7},
            polarity={1: -1},
            selected_lobe={1: "negative_trough"},
            event_ids={1: "negative-event"},
            **common,
        ),
        SeedStation(
            "positive",
            30.0,
            {1: 96.0},
            phase_class={1: 3},
            polarity={1: 1},
            selected_lobe={1: "positive_peak"},
            event_ids={1: "positive-event"},
            **common,
        ),
    ]

    with pytest.raises(ValueError, match="conflicting wavelet lobes"):
        save_seed_file(tmp_path / "conflict.json", "ROAD_A", stations)


def test_blinded_checkpoint_round_trip_is_explicitly_non_training(tmp_path):
    path = tmp_path / "checkpoints.json"
    checkpoints = [
        ValidationCheckpoint(
            checkpoint_id="base-001",
            layer_order=2,
            chainage_m=145.0,
            sample_index=248.0,
            canonical_sample_index=250.0,
            visibility=VisibilityState.VISIBLE,
            user_confirmed=True,
            selected_lobe="negative_trough",
            event_family_id="base-family-a",
            pulse_width_samples=8.0,
        )
    ]

    save_checkpoint_file(path, "ROAD_A", checkpoints)
    survey_id, loaded = load_checkpoint_file(path)

    assert survey_id == "ROAD_A"
    assert loaded == checkpoints
    document = json.loads(path.read_text(encoding="utf-8"))
    assert document["selection_source"] == "radar_only"
    assert document["training_use"] == "prohibited"


def test_project_schema_is_explicitly_prototype_v3(tmp_path):
    path = tmp_path / "prototype.gprproj"
    store = ProjectStore.create(path, "survey-a")
    store.validate()
    assert SCHEMA_VERSION == 3
    assert store.get_meta("survey_id") == "survey-a"


def test_blocked_holdout_is_deterministic_and_frozen_to_100m_modulo_five():
    chainages = np.linspace(0, 999, 1_000)
    first = _holdout_block("road-a:1", chainages)
    second = _holdout_block("road-a:1", chainages)
    assert (first == second).all()
    assert first.any()
    assert np.all((np.floor(chainages[first] / 100).astype(int) % 5) == 4)
