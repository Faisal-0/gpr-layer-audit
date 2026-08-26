from __future__ import annotations

import json

from openpyxl import load_workbook

from gpr_layer_audit.design import compare_with_design
from gpr_layer_audit.export import export_audit_package
from gpr_layer_audit.models import AcquisitionFileSet, DesignSegment, LayerDesign
from gpr_layer_audit.processing import AnalysisOptions, analyze_acquisition
from gpr_layer_audit.project import ProjectStore


def _result(synthetic_acquisition):
    road_path, plate_path, _ = synthetic_acquisition
    return analyze_acquisition(
        AcquisitionFileSet(road_path),
        AcquisitionFileSet(plate_path),
        AnalysisOptions(stack_size=4, accept_scan_dielectric=True),
    )


def test_design_comparison_does_not_change_measurement(synthetic_acquisition):
    result = _result(synthetic_acquisition)
    original = [(item.bottom_sample, item.thickness_mm) for item in result.thickness]
    segments = [DesignSegment("road", 0, 1000, "Asphalt", 60, -5, 5)]
    compared = compare_with_design(result.thickness, segments)
    assert [(item.bottom_sample, item.thickness_mm) for item in compared] == original
    asphalt = [item for item in compared if item.layer_name == "Asphalt"]
    assert asphalt and all(item.design_thickness_mm == 60 for item in asphalt)


def test_project_store_persists_anchor_and_results(tmp_path, synthetic_acquisition):
    result = _result(synthetic_acquisition)
    store = ProjectStore.create(tmp_path / "test.gprproj", "Synthetic")
    store.set_file("road", result.source.dzt_path)
    assert store.file_paths()["road"] == result.source.dzt_path
    designs = [
        LayerDesign(1, "Asphalt", 50.8, 7.0),
        LayerDesign(2, "Base course", 101.6, 7.0),
        LayerDesign(3, "Sub-base course", None, 7.0),
    ]
    store.set_layer_designs(designs)
    assert store.layer_designs() == designs
    store.add_anchor(1, 12.5, 98)
    assert store.anchors() == {1: [(12.5, 98.0)]}
    assert store.remove_last_anchor() == (1, 12.5, 98.0)
    assert store.anchors() == {}
    run_id = store.save_analysis(result)
    assert run_id == 1


def test_audit_export_contains_expected_artifacts(tmp_path, synthetic_acquisition):
    result = _result(synthetic_acquisition)
    package = export_audit_package(result, tmp_path)
    expected = {
        "gpr_layer_audit.xlsx",
        "thickness_results.csv",
        "thickness_results.geojson",
        "annotated_radargram.png",
        "interface_observations.csv",
        "layer_profiles.csv",
        "layer_profiles.png",
        "manifest.json",
    }
    assert expected.issubset({path.name for path in package.iterdir()})
    assert package.with_suffix(".zip").exists()
    workbook = load_workbook(package / "gpr_layer_audit.xlsx", read_only=True)
    assert {
        "Summary",
        "Thickness Results",
        "Exceptions",
        "Interface Picks",
        "Layer Profiles",
        "Structural Anomalies",
        "Candidate Events",
    }.issubset(workbook.sheetnames)
    manifest = json.loads((package / "manifest.json").read_text(encoding="utf-8"))
    assert manifest["trace_bins"] == 60
